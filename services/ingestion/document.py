import hashlib
import zipfile
from dataclasses import dataclass
from pathlib import Path

from ml.text_normalization.service import normalize_text


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    source_id: str
    filename: str
    content_type: str
    text: str
    page_count: int | None
    extraction_status: str
    file_sha256: str | None = None
    pages: tuple[ExtractedPage, ...] = ()
    error: str | None = None


def _file_sha256(file: Path) -> str:
    digest = hashlib.sha256()
    with file.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_document(path: str, source_id: str) -> ExtractedDocument:
    file = Path(path)
    suffix = file.suffix.lower()
    try:
        file_sha256 = _file_sha256(file)
        if suffix == ".pdf":
            import fitz

            with fitz.open(file) as doc:
                pages = tuple(
                    ExtractedPage(page_number=index + 1, text=normalize_text(page.get_text("text")))
                    for index, page in enumerate(doc)
                )
            return ExtractedDocument(
                source_id, file.name, "application/pdf",
                "\n\n".join(page.text for page in pages),
                len(pages), "success", file_sha256, pages,
            )
        if suffix == ".docx":
            from docx import Document

            doc = Document(file)
            text = normalize_text("\n".join(p.text for p in doc.paragraphs if p.text.strip()))
            page = ExtractedPage(page_number=1, text=text)
            return ExtractedDocument(
                source_id, file.name,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                text, 1, "success", file_sha256, (page,),
            )
        return ExtractedDocument(
            source_id, file.name, "application/octet-stream", "", None,
            "unsupported_type", file_sha256, (), f"Unsupported file type: {suffix or 'none'}",
        )
    except (OSError, RuntimeError, ValueError, zipfile.BadZipFile) as exc:
        return ExtractedDocument(
            source_id, file.name, "application/octet-stream", "", None,
            "extraction_failed", None, (), str(exc),
        )
