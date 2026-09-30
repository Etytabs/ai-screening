from tempfile import NamedTemporaryFile

import fitz
from docx import Document

from services.screening.pipeline import screen_document


def test_screening_rejects_unsupported_documents():
    with NamedTemporaryFile(suffix=".txt") as tmp:
        tmp.write(b"demo")
        tmp.flush()
        result = screen_document(tmp.name, "TEST-001")
    assert result["status"] == "unsupported_type"
    assert result["extraction"]["extraction_status"] == "unsupported_type"


def test_screening_processes_pdf_document():
    with NamedTemporaryFile(suffix=".pdf") as tmp:
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), "Abstract\nMethodology\nBudget")
        document.save(tmp.name)
        document.close()

        result = screen_document(tmp.name, "TEST-PDF")

    assert result["status"] == "screened"
    assert result["extraction"]["extraction_status"] == "success"
    assert result["extraction"]["page_count"] == 1
    assert result["completeness"]["total"] == 6
    assert result["human_review_required"] is True


def test_screening_processes_docx_document():
    with NamedTemporaryFile(suffix=".docx") as tmp:
        document = Document()
        document.add_paragraph("Abstract")
        document.add_paragraph("Methodology")
        document.add_paragraph("Budget")
        document.save(tmp.name)

        result = screen_document(tmp.name, "TEST-DOCX")

    assert result["status"] == "screened"
    assert result["extraction"]["extraction_status"] == "success"
    assert result["completeness"]["total"] == 6
    assert result["human_review_required"] is True
